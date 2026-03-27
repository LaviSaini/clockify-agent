"""Write Markdown and Word reports under reports/."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document


def project_root() -> Path:
    return Path(__file__).resolve().parent.parent


def reports_dir() -> Path:
    d = project_root() / "reports"
    d.mkdir(parents=True, exist_ok=True)
    return d


def write_markdown(title: str, body: str) -> Path:
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%SZ")
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in title)[:80]
    path = reports_dir() / f"{ts}_{safe}.md"
    path.write_text(f"# {title}\n\n{body}\n", encoding="utf-8")
    return path


def write_docx(title: str, body: str) -> Path:
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%SZ")
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in title)[:80]
    path = reports_dir() / f"{ts}_{safe}.docx"
    doc = Document()
    doc.add_heading(title, level=1)
    for block in body.split("\n\n"):
        doc.add_paragraph(block.strip())
    doc.save(str(path))
    return path
