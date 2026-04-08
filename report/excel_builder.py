from __future__ import annotations

import os
from collections import defaultdict
from datetime import datetime, timedelta

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

RED = PatternFill("solid", fgColor="FFCCCC")
AMBER = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


def _append_header_row(ws, values: list) -> None:
    row = []
    for v in values:
        c = WriteOnlyCell(ws, v)
        c.font = BOLD
        row.append(c)
    ws.append(row)


def _append_filled_row(ws, values: list, fill: PatternFill) -> None:
    row = []
    for v in values:
        c = WriteOnlyCell(ws, v)
        c.fill = fill
        row.append(c)
    ws.append(row)


def _append_write_row(ws, values: list) -> None:
    """Write-only sheets: append a full row using WriteOnlyCell (matches header pattern)."""
    row = [WriteOnlyCell(ws, v) for v in values]
    ws.append(row)


def _name_for_id(data: dict, uid: str) -> str:
    return ((data.get("user_name_by_id") or {}).get(uid, "") or "").strip()


def _email_for_id(data: dict, uid: str) -> str:
    return ((data.get("user_email_by_id") or {}).get(uid, "") or "").strip()


def _total_hours_for_id(data: dict, uid: str) -> float:
    raw = (data.get("total_hours_by_user_id") or {}).get(uid)
    if raw is None:
        return 0.0
    try:
        return float(raw)
    except (TypeError, ValueError):
        return 0.0


def _leave_days_for_id(data: dict, uid: str) -> int:
    """Count of calendar days on approved leave overlapping the report range (from HR CSV)."""
    sk = str(uid).strip()
    days_map = data.get("leave_days_by_user_id") or {}
    raw = days_map.get(sk)
    if raw is not None:
        try:
            return int(raw)
        except (TypeError, ValueError):
            pass
    dates_map = data.get("leave_dates_by_user_id") or {}
    dates = dates_map.get(sk)
    if dates is not None:
        return len(dates)
    return 0


def _row_uid(row: dict) -> str:
    return ((row.get("user_id") or row.get("user") or "") or "").strip()


def _weekday_count_in_range(start_date: str, end_date: str) -> int:
    """Mon–Fri days inclusive between start and end (same basis as detect_missing_logs)."""
    start = datetime.strptime(start_date, "%Y-%m-%d")
    end = datetime.strptime(end_date, "%Y-%m-%d")
    n = 0
    current = start
    while current <= end:
        if current.weekday() < 5:
            n += 1
        current += timedelta(days=1)
    return n


def _deficit_exceeded_hours_for_day(data: dict, row: dict) -> float:
    """Signed vs daily target: logged − MIN_HOURS (negative = deficit, positive = exceeded)."""
    try:
        target = float(data.get("min_hours_per_day") or 8)
        logged = float(row.get("hours_logged") or 0)
        return round(logged - target, 2)
    except (TypeError, ValueError):
        return 0.0


def _total_deficit_exceeded_hours(
    data: dict, uid: str, start_date: str, end_date: str
) -> float:
    """
    Period total: actual hours − (MIN_HOURS × weekday count), e.g. 37 − 40 = −3, 45 − 40 = +5.
    """
    min_h = float(data.get("min_hours_per_day") or 8)
    expected = min_h * _weekday_count_in_range(start_date, end_date)
    actual = _total_hours_for_id(data, uid)
    return round(actual - expected, 2)


def _project_hours_rows_sorted_by_email(data: dict) -> list[tuple[str, str, str, float]]:
    """
    Aggregate hours by (user, project) and sort by email, then project.
    Returns tuples: (user_name, email, project, hours).
    """
    by_user_project: dict[tuple[str, str], float] = defaultdict(float)
    entries_by_user = data.get("time_entries_by_user") or {}
    for uid, entries in entries_by_user.items():
        uid_s = str(uid).strip()
        for entry in entries:
            project = (entry.get("project") or "").strip() or "Unassigned"
            try:
                hours = float(entry.get("hours") or 0)
            except (TypeError, ValueError):
                hours = 0.0
            by_user_project[(uid_s, project)] += hours

    rows: list[tuple[str, str, str, float]] = []
    for (uid, project), total in by_user_project.items():
        rows.append(
            (
                _name_for_id(data, uid),
                _email_for_id(data, uid),
                project,
                round(total, 2),
            )
        )

    rows.sort(key=lambda r: ((r[1] or "").strip().lower(), (r[2] or "").strip().lower()))
    return rows


def build_excel_report(data: dict, start_date: str, end_date: str) -> str:
    wb = openpyxl.Workbook(write_only=True)

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.create_sheet("Summary")
    _append_header_row(
        ws1,
        [
            "User",
            "Email",
            "Total Hours",
            "Leave days",
            "Missing Days",
            "Incomplete Days",
            "Total Deficit / Exceeded Hours",
            "Flagged Entries",
        ],
    )

    missing_count = defaultdict(lambda: {"Missing": 0, "Incomplete": 0})
    for row in data.get("missing_logs", []):
        rid = _row_uid(row)
        if rid:
            missing_count[rid][row["severity"]] += 1

    desc_count = defaultdict(int)
    for row in data.get("poor_descriptions", []):
        rid = _row_uid(row)
        if rid:
            desc_count[rid] += 1

    users_for_summary = data.get("workspace_user_ids")
    if not users_for_summary:
        issue_ids = set(missing_count.keys()) | set(desc_count.keys())
        users_for_summary = sorted(issue_ids, key=lambda x: x or "")

    for uid in users_for_summary:
        _append_write_row(
            ws1,
            [
                _name_for_id(data, uid),
                _email_for_id(data, uid),
                _total_hours_for_id(data, uid),
                _leave_days_for_id(data, uid),
                missing_count[uid]["Missing"],
                missing_count[uid]["Incomplete"],
                _total_deficit_exceeded_hours(data, uid, start_date, end_date),
                desc_count[uid],
            ],
        )

    # ── Sheet 2: Missing Logs ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Missing Logs")
    _append_header_row(
        ws2,
        ["User", "Email", "Date", "Day", "Hours Logged", "Deficit / Exceeded Hours", "Severity"],
    )

    for row in data.get("missing_logs", []):
        fill = RED if row["severity"] == "Missing" else AMBER
        uid = (row.get("user_id") or "").strip()
        name_cell = _name_for_id(data, uid) if uid else (row.get("user") or "").strip()
        email_cell = _email_for_id(data, uid) if uid else ""
        deficit = _deficit_exceeded_hours_for_day(data, row)
        _append_filled_row(
            ws2,
            [
                name_cell,
                email_cell,
                row["date"],
                row["day"],
                row["hours_logged"],
                deficit,
                row["severity"],
            ],
            fill,
        )

    # ── Sheet 3: Poor Descriptions ────────────────────────────────────────────
    ws3 = wb.create_sheet("Poor Descriptions")
    _append_header_row(
        ws3, ["User", "Email", "Date", "Project", "Description", "Score", "Reason"]
    )

    for row in data.get("poor_descriptions", []):
        fill = RED if row["score"] == 1 else AMBER
        uid = (row.get("user_id") or "").strip()
        name_cell = _name_for_id(data, uid) if uid else (row.get("user") or "").strip()
        email_cell = _email_for_id(data, uid) if uid else ""
        _append_filled_row(
            ws3,
            [
                name_cell,
                email_cell,
                row["date"],
                row.get("project", ""),
                row["description"],
                row["score"],
                row["reason"],
            ],
            fill,
        )

    # ── Sheet 4: Project Hours by Employee ────────────────────────────────────
    ws4 = wb.create_sheet("Project Hours by Employee")
    _append_header_row(ws4, ["User", "Email", "Project", "Hours"])
    for user, email, project, hours in _project_hours_rows_sorted_by_email(data):
        _append_write_row(ws4, [user, email, project, hours])

    # ── Save ──────────────────────────────────────────────────────────────────
    os.makedirs("output", exist_ok=True)
    filename = f"output/loglens_report_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
    wb.save(filename)
    print(f"Report saved -> {filename}")
    return filename
